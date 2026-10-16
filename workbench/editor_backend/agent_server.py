"""Local Agent API plus compatibility with the original Vivid editor."""
from pathlib import Path
from contextlib import asynccontextmanager
import argparse, asyncio, base64, csv, io, json, secrets, sys, time, webbrowser
from urllib.parse import urlsplit
sys.path.insert(0, str(Path(__file__).resolve().parent))
from fastapi import FastAPI, Request
from fastapi.responses import JSONResponse, FileResponse, StreamingResponse
from ag_ui.encoder import EventEncoder
from ag_ui import core as events
from schema import SPECS, sample, prepare, Invalid, validate_language_settings
import library, object_editor
from server import job_folder, FRONTEND_FILES
from agent import store, tools, jobs, provider, runtime

WEB = Path(__file__).resolve().parent.parent
TOKEN = secrets.token_urlsafe(32)

@asynccontextmanager
async def lifespan(app):
    with store.instance_lock():
        runtime.init()
        yield
        for rid in list(runtime.THREADS): runtime.stop(rid)

app = FastAPI(lifespan=lifespan, docs_url=None, redoc_url=None, openapi_url=None)

def failure(message, code=400): return JSONResponse({'error':{'message':str(message)}}, status_code=code)

@app.middleware('http')
async def local_only(request, call_next):
    host = urlsplit('http://' + request.headers.get('host','')).hostname
    if host not in ['127.0.0.1','localhost','::1']: return failure('仅接受本机访问。',403)
    origin = request.headers.get('origin')
    if origin and origin != str(request.base_url).rstrip('/'): return failure('请求来源不匹配。',403)
    if request.method not in ['GET','HEAD']:
        if 'application/json' not in request.headers.get('content-type',''): return failure('请求需要 JSON。',415)
        if int(request.headers.get('content-length','0')) > 20*1024*1024: return failure('请求最大 20 MB。',413)
        if request.url.path.startswith('/api/agent/') and request.headers.get('x-vivid-token') != TOKEN: return failure('页面已过期，请刷新后重试。',403)
    response = await call_next(request)
    response.headers['X-Content-Type-Options']='nosniff'
    response.headers['Referrer-Policy']='same-origin'
    response.headers['Cache-Control']='no-store' if request.url.path.startswith('/api/') else 'no-cache'
    return response

@app.exception_handler(Invalid)
async def invalid(request, error): return JSONResponse({'error':error.detail}, status_code=400)
@app.exception_handler(store.Conflict)
async def conflict(request, error): return failure(error,409)
@app.exception_handler(ValueError)
@app.exception_handler(KeyError)
@app.exception_handler(TypeError)
async def bad_input(request, error): return failure(error)
@app.exception_handler(FileNotFoundError)
async def missing_file(request, error): return failure('本机记录或文件不存在。',404)

def valid_project(p):
    if not isinstance(p,dict) or p.get('template') not in SPECS: raise ValueError('模板不存在。')
    validate_language_settings(p.get('settings',{}))
    for table in (p.get('tables') or {'data':p.get('table')}).values(): tools.validate_table(table)
    object_editor.validate(p)
    return p

def import_file(data):
    name=data.get('name','');raw=base64.b64decode(data.get('content',''),validate=True)
    if len(raw)>12*1024*1024: raise ValueError('数据文件最大 12 MB。')
    if name.lower().endswith('.xlsx'):
        import openpyxl
        workbook=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        try:
            sheets=workbook.sheetnames;sheet=data.get('sheet') or sheets[0]
            if data.get('list_sheets'):return {'table':None,'sheets':sheets,'sheet':None}
            if sheet not in sheets: raise ValueError('工作表不存在。')
            grid=[]
            for row in workbook[sheet].iter_rows(values_only=True):
                if len(row)>256:raise ValueError('最多支持 256 列。')
                grid.append([v.isoformat() if hasattr(v,'isoformat') else v for v in row])
                if len(grid)>20001:raise ValueError('最多支持 20000 行数据。')
        finally:workbook.close()
    elif name.lower().endswith(('.csv','.tsv')):
        try:text=raw.decode('utf-8-sig')
        except UnicodeDecodeError:text=raw.decode('gb18030')
        try:dialect=csv.Sniffer().sniff(text[:8192],delimiters=',;\t')
        except csv.Error:dialect=csv.excel_tab if name.lower().endswith('.tsv') else csv.excel
        grid=[];sheets=[];sheet=None
        for row in csv.reader(io.StringIO(text),dialect):
            grid.append(row)
            if len(grid)>20001:raise ValueError('最多支持 20000 行数据。')
    else:raise ValueError('请选择 CSV、TSV 或 XLSX 文件。')
    while grid and all(v in ['',None] for v in grid[-1]):grid.pop()
    if len(grid)<2:raise ValueError('文件需要表头和至少一行数据。')
    width=len(grid[0])
    while width and all(len(row)<width or row[width-1] in ['',None] for row in grid):width-=1
    if any(len(row)>width and any(v not in ['',None] for v in row[width:]) for row in grid):raise ValueError('部分数据行比表头更宽。')
    table={'headers':[str('' if v is None else v).strip() for v in grid[0][:width]],'rows':[[('' if v is None else v) for v in (row[:width]+['']*max(0,width-len(row)))] for row in grid[1:]]}
    tools.validate_table(table)
    return {'table':table,'sheets':sheets,'sheet':sheet}

@app.get('/api/health')
def health():return {'app':'vivid-workbench','root':str(WEB),'version':3,'agent':True}
@app.get('/api/bootstrap')
def bootstrap():return {'token':TOKEN,'model':provider.config(public=True)}
@app.get('/api/templates')
def templates():return {'templates':list(SPECS.values()),'palettes':json.loads((WEB/'editor_backend/_utils/palettes.json').read_text(encoding='utf8'))['palettes']}
@app.get('/api/sample')
def get_sample(id:str):return sample(id)
@app.get('/api/library')
def get_library(kind:str='projects',id:str=''):return library.get(kind,id) if id else {'items':library.listing(kind)}
@app.post('/api/library')
def save_library(data:dict):return library.save(data.get('kind','projects'),data)
@app.post('/api/import')
def file_import(data:dict):return import_file(data)
@app.post('/api/validate')
def validate(data:dict):
    object_editor.validate(data);_,notes=prepare(data);return {'ok':True,'notes':notes}
@app.post('/api/render',status_code=202)
def render(data:dict):return jobs.submit(data)
@app.get('/api/jobs/{job}')
def job_status(job:str):
    jobs.folder(job)
    if (jobs.JOBS/job).exists():return jobs.status(job)
    f=job_folder(job)
    if not f.exists():return failure('绘图任务不存在。',404)
    result=json.loads((f/'status.json').read_text(encoding='utf8'))
    if result['state']=='done':result['base']=f'/workspace/jobs/{job}/'
    return result

@app.get('/api/agent/catalog')
def catalog():return tools.catalog()
@app.get('/api/agent/templates/{tid}')
def template_contract(tid:str):return tools.template_info(tid)
@app.get('/api/agent/sessions')
def sessions():return {'items':store.sessions()}
@app.post('/api/agent/sessions')
def new_session(data:dict):return store.create_session(data.get('name','新的绘图项目'))
@app.get('/api/agent/sessions/{sid}')
def session(sid:str):return store.session(sid)
@app.post('/api/agent/sessions/{sid}/datasets')
def dataset(sid:str,data:dict):return store.dataset(sid,data.get('name','数据表'),tools.validate_table(data['table']))
@app.post('/api/agent/sessions/{sid}/charts')
def chart(sid:str,data:dict):
    p=data.get('project') or sample(data['template'])
    if 'project' not in data:p['name'] += ' · 示例'
    return store.create_chart(sid,valid_project(p),'载入模板示例' if 'project' not in data else '导入图表')
@app.post('/api/agent/sessions/{sid}/bind')
def bind_chart(sid:str,data:dict):
    workspace=store.session(sid)
    p=tools.bind_template(data['template_id'],workspace['datasets'],data['bindings'],data.get('mapping'),data.get('name'),data.get('settings'))
    return store.create_chart(sid,valid_project(p),'使用工作区数据创建')
@app.get('/api/agent/charts/{cid}')
def chart_get(cid:str,rev:int|None=None):return store.chart(cid,rev)
@app.get('/api/agent/charts/{cid}/revisions')
def revisions(cid:str):return {'items':store.revisions(cid)}
@app.post('/api/agent/charts/{cid}/save')
def save(cid:str,data:dict):
    p=valid_project(data['project']);job=data.get('job')
    if job:
        if jobs.status(job)['state']!='done' or json.loads((jobs.folder(job)/'project.json').read_text(encoding='utf8'))!=p:raise ValueError('预览与本次项目版本不一致。')
    return store.save_revision(cid,data['base_revision'],p,data.get('label','手动编辑'),'user',job)
@app.post('/api/agent/charts/{cid}/restore')
def restore(cid:str,data:dict):
    old=store.chart(cid,data['revision'])
    return store.save_revision(cid,data['base_revision'],old['project'],'恢复版本 '+str(data['revision']),'user',old['job'])
@app.post('/api/agent/charts/{cid}/branch')
def branch(cid:str,data:dict):
    old=store.chart(cid,data.get('revision'));p=old['project'];p['name']+=' · 副本'
    return store.create_chart(old['session'],p,'从版本 '+str(old['revision'])+' 创建分支','user',None)
@app.post('/api/agent/sessions/{sid}/runs')
def run_start(sid:str,data:dict):return runtime.start(sid,data['prompt'],data.get('chart_id'),data.get('selected_object'),data.get('target'),data.get('dataset_ids'))
@app.get('/api/agent/runs/{rid}')
def run_get(rid:str):return store.run(rid)
@app.post('/api/agent/runs/{rid}/stop')
def stop(rid:str,data:dict):return runtime.stop(rid)
@app.post('/api/agent/runs/{rid}/resume')
def resume(rid:str,data:dict):return runtime.resume(rid,data.get('response'))
@app.post('/api/agent/runs/{rid}/candidate')
def candidate(rid:str,data:dict):
    r=store.run(rid);result=r['body'].get('result',{})
    if not result.get('conflict'):raise ValueError('没有待保留的候选图。')
    c=store.create_chart(r['session'],result['project'],'保留 Agent 候选图','user',result.get('job'),operation=rid+':candidate')
    return c

@app.get('/api/agent/runs/{rid}/events')
async def stream(rid:str,request:Request,after:int=0):
    store.run(rid);after=max(after,int(request.headers.get('last-event-id','0')))
    models={getattr(events,name).model_fields['type'].default:getattr(events,name) for name in ['RunStartedEvent','RunFinishedEvent','RunErrorEvent','StepStartedEvent','StepFinishedEvent','CustomEvent','StateSnapshotEvent']}
    encoder=EventEncoder()
    async def generate():
        cursor=after;idle=0
        while not await request.is_disconnected():
            rows=store.events(rid,cursor)
            for row in rows:
                event=row['event'];cls=models.get(event['type'])
                if cls is None:raise RuntimeError('未知事件类型')
                yield 'id: '+str(row['seq'])+'\n'+encoder.encode(cls.model_validate(event));cursor=row['seq']
            status=store.run(rid)['status']
            if not rows and status in ['done','error','cancelled','waiting','interrupted']:
                yield 'event: settled\ndata: '+json.dumps({'status':status})+'\n\n';break
            idle+=1
            if idle%30==0:yield ': keepalive\n\n'
            await asyncio.sleep(.3)
    return StreamingResponse(generate(),media_type='text/event-stream',headers={'X-Accel-Buffering':'no'})

@app.get('/api/agent/config')
def config():return provider.config(public=True)
@app.post('/api/agent/config')
def save_config(data:dict):return provider.save_config(data)
@app.post('/api/agent/config/test')
def test_config(data:dict):
    schema={'type':'object','properties':{'message':{'type':'string'}},'required':['message'],'additionalProperties':False}
    result,usage=provider.call('请只返回 message=连接成功，不调用任何工具。',schema,store.ROOT/'connection-test')
    return {'ok':True,'message':result['message'],'usage':usage}

@app.get('/{path:path}')
def static(path:str):
    path=path or 'index.html';parts=Path(path).parts
    if path in FRONTEND_FILES:path='frontend/'+path
    if path.startswith('workspace/jobs/'):
        chunks=path.split('/');job=chunks[2];jobs.folder(job);base=job_folder(job).resolve();target=(base/('/'.join(chunks[3:]))).resolve()
        if not target.is_relative_to(base) or target.suffix.lower() not in ['.png','.jpg','.svg','.pdf','.json','.csv','.py','.zip','.html','.js']:return failure('文件不可访问。',404)
    elif path in ['index.html','editor.html','agent.html'] or path.startswith(('assets/','frontend/')):
        base=(WEB/path.split('/')[0]).resolve() if '/' in path else WEB.resolve();target=(WEB/path).resolve()
        if not target.is_relative_to(base):return failure('文件不可访问。',404)
    else:return failure('文件不可访问。',404)
    if not target.is_file():return failure('文件不存在。',404)
    return FileResponse(target)

if __name__=='__main__':
    import uvicorn, threading
    parser=argparse.ArgumentParser();parser.add_argument('--port',type=int,default=8766);parser.add_argument('--open',action='store_true');args=parser.parse_args()
    import httpx
    try:
        running=httpx.get(f'http://127.0.0.1:{args.port}/api/health',timeout=1,trust_env=False).json()
    except (httpx.HTTPError,ValueError):running=None
    if running:
        if running.get('agent') and running.get('root')==str(WEB):
            if args.open:webbrowser.open(f'http://127.0.0.1:{args.port}/index.html')
            sys.exit(0)
        raise SystemExit(f'{args.port} 端口已有其他服务。请关闭旧工作台，或使用 --port 8767。')
    if args.open:threading.Timer(1,lambda:webbrowser.open(f'http://127.0.0.1:{args.port}/index.html')).start()
    uvicorn.run(app,host='127.0.0.1',port=args.port,log_level='warning')
