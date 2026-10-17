"""Loopback-only local UI and Python renderer. Start with --open for the browser."""
from pathlib import Path
import os,sys,json,uuid,subprocess,threading,base64,io,csv,urllib.parse,urllib.request,webbrowser,shutil,socket
from http.server import ThreadingHTTPServer,SimpleHTTPRequestHandler
ROOT=Path(__file__).resolve().parent;WEB=ROOT.parent;JOBS=WEB/'workspace/jobs'
ARCHIVE_REGISTRY=WEB/'workspace/archive-jobs.json'
FRONTEND_FILES=frozenset(['app.js','style.css','editor.js','editor.css','grid.js',
    'interactions.js','interactions.css','direct-edit.js','data-transform.js','studio.js','studio.css'])
sys.path.insert(0,str(ROOT))
from schema import SPECS,sample,defaults,prepare,Invalid
import object_editor,library
from architecture_assets import architecture_file
PORT=8766
def python_exe():
    exe=Path(sys.executable)
    return str(exe.with_name('python.exe')) if exe.name.lower()=='pythonw.exe' else str(exe)
semaphore=threading.Semaphore(2)

def job_folder(job):
    """Keep old job URLs usable after moving completed jobs to the local archive."""
    live=JOBS/job
    if live.is_dir() or not ARCHIVE_REGISTRY.is_file():return live
    archive=(WEB.parent/'归档').resolve()
    for relative in json.loads(ARCHIVE_REGISTRY.read_text(encoding='utf8')).get('roots',[]):
        root=(WEB/relative).resolve()
        if not root.is_relative_to(archive):continue
        folder=(root/job).resolve()
        if folder.is_relative_to(root) and folder.is_dir():return folder
    return live
def execute(folder):
    with semaphore:
        try:
            with (folder/'render.log').open('w',encoding='utf8') as log:
                p=subprocess.run([python_exe(),'-X','utf8',str(ROOT/'worker.py'),str(folder)],cwd=folder,stdout=log,stderr=log,timeout=180,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
            if p.returncode and not (folder/'status.json').exists():raise RuntimeError('绘图进程退出。')
        except Exception as e:(folder/'status.json').write_text(json.dumps({'state':'error','message':str(e)},ensure_ascii=False),encoding='utf8')

class Handler(SimpleHTTPRequestHandler):
    def __init__(self,*a,**k):super().__init__(*a,directory=str(WEB),**k)
    def log_message(self,*a):pass
    def respond(self,data,code=200):
        content=json.dumps(data,ensure_ascii=False,allow_nan=False).encode('utf8')
        self.send_response(code);self.send_header('Content-Type','application/json; charset=utf-8');self.send_header('Content-Length',str(len(content)));self.send_header('Cache-Control','no-store');self.end_headers();self.wfile.write(content)
    def end_headers(self):
        path=urllib.parse.urlsplit(self.path).path
        if path in ['/', '/index.html', '/architecture.html', '/architecture-editor.html', '/architecture-agent.html', '/editor.html', '/assets/catalog.js', '/assets/editable.js'] or path.startswith(('/frontend/', '/architecture-assets/')) or path.lstrip('/') in FRONTEND_FILES:
            self.send_header('Cache-Control','no-cache')
        self.send_header('X-Content-Type-Options','nosniff');super().end_headers()
    def translate_path(self,path):
        decoded=urllib.parse.unquote(urllib.parse.urlsplit(path).path)
        if decoded.startswith('/architecture-assets/'):
            try:return str(architecture_file(decoded[len('/architecture-assets/'):]))
            except (FileNotFoundError, OSError, ValueError):return str(WEB/'__missing_architecture_resource__')
        if decoded.startswith('/') and decoded[1:] in FRONTEND_FILES:
            return str(WEB/'frontend'/decoded[1:])
        parts=decoded.split('/')
        if len(parts)>=4 and parts[1:3]==['workspace','jobs']:
            job=parts[3]
            if len(job)==32 and all(c in '0123456789abcdef' for c in job):
                folder=job_folder(job).resolve()
                target=folder.joinpath(*parts[4:]).resolve()
                if target.is_relative_to(folder):return str(target)
        return super().translate_path(path)
    def do_GET(self):
        url=urllib.parse.urlsplit(self.path);path=url.path;query=urllib.parse.parse_qs(url.query)
        try:
            if path=='/api/health':return self.respond({'app':'vivid-workbench','root':str(WEB),'version':3})
            if path=='/api/templates':
                palettes=json.loads((ROOT/'_utils/palettes.json').read_text(encoding='utf8'))['palettes']
                return self.respond({'templates':list(SPECS.values()),'palettes':palettes})
            if path=='/api/sample':
                tid=query.get('id',[''])[0]
                if tid not in SPECS:raise Invalid('未知模板。')
                return self.respond(sample(tid))
            if path=='/api/library':
                kind=query.get('kind',['projects'])[0];key=query.get('id',[''])[0]
                return self.respond(library.get(kind,key) if key else {'items':library.listing(kind)})
            if path.startswith('/api/jobs/'):
                job=path.rsplit('/',1)[-1]
                if len(job)!=32 or any(c not in '0123456789abcdef' for c in job):raise Invalid('无效任务编号。')
                folder=job_folder(job)
                if not folder.exists():return self.respond({'error':{'message':'任务不存在。'}},404)
                result=json.loads((folder/'status.json').read_text(encoding='utf8'))
                if result['state']=='done':result['base']=f'/workspace/jobs/{job}/'
                return self.respond(result)
            if path.startswith('/api/'):return self.respond({'error':{'message':'接口不存在。'}},404)
            return super().do_GET()
        except Invalid as e:return self.respond({'error':e.detail},400)
        except Exception as e:return self.respond({'error':{'message':str(e)}},500)
    def do_POST(self):
        origin=self.headers.get('Origin')
        if origin and origin not in [f'http://127.0.0.1:{PORT}',f'http://localhost:{PORT}']:return self.respond({'error':{'message':'仅接受本地工作台请求。'}},403)
        try:
            if 'application/json' not in self.headers.get('Content-Type',''):raise Invalid('请求需要 JSON 数据。')
            length=int(self.headers.get('Content-Length','0'))
            if not 0<length<=20*1024*1024:raise Invalid('请求大小应在 20 MB 以内。')
            data=json.loads(self.rfile.read(length))
            if self.path=='/api/library':
                return self.respond(library.save(data.get('kind','projects'),data))
            if self.path=='/api/render':
                object_editor.validate(data)
                _,notes=prepare(data)
                # Source comes only from the shipped catalog, never from request code.
                job=uuid.uuid4().hex;folder=JOBS/job;folder.mkdir(parents=True)
                data=dict(data);data['source_version']='workbench-editor-2'
                (folder/'project.json').write_text(json.dumps(data,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf8')
                source=WEB/'assets/sources'/data['template']
                shutil.copy2(source/'preview.py',folder/'source.py');shutil.copy2(source/'original.py',folder/'original.py')
                (folder/'status.json').write_text('{"state":"queued","message":"等待绘图"}',encoding='utf8')
                threading.Thread(target=execute,args=(folder,),daemon=True).start()
                return self.respond({'job':job,'notes':notes},202)
            if self.path=='/api/validate':
                object_editor.validate(data)
                _,notes=prepare(data);return self.respond({'ok':True,'notes':notes})
            if self.path=='/api/import':
                name=data.get('name','');raw=base64.b64decode(data.get('content',''),validate=True)
                if len(raw)>12*1024*1024:raise Invalid('数据文件最大 12 MB。')
                if name.lower().endswith('.xlsx'):
                    import openpyxl
                    workbook=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
                    sheets=workbook.sheetnames;sheet=data.get('sheet') or sheets[0]
                    if sheet not in sheets:raise Invalid('工作表不存在。')
                    grid=[]
                    for row in workbook[sheet].iter_rows(values_only=True):
                        grid.append([v.isoformat() if hasattr(v,'isoformat') else v for v in row])
                        if len(grid)>20001:raise Invalid('最多支持 20000 行数据。')
                    workbook.close()
                elif name.lower().endswith(('.csv','.tsv')):
                    try:text=raw.decode('utf-8-sig')
                    except UnicodeDecodeError:text=raw.decode('gb18030')
                    try:dialect=csv.Sniffer().sniff(text[:8192],delimiters=',;\t')
                    except csv.Error:dialect=csv.excel_tab if name.lower().endswith('.tsv') else csv.excel
                    grid=list(csv.reader(io.StringIO(text),dialect));sheets=[];sheet=None
                else:raise Invalid('请选择 CSV、TSV 或 XLSX 文件。')
                while grid and all(v in ['',None] for v in grid[-1]):grid.pop()
                if len(grid)<2:raise Invalid('文件需要表头和至少一行数据。')
                width=len(grid[0])
                while width and all(len(row)<width or row[width-1] in ['',None] for row in grid):width-=1
                if not 1<=width<=256 or len(grid)>20001 or width*(len(grid)-1)>500000:raise Invalid('文件需为 1–256 列、最多 20000 行，总单元格不超过 500000。')
                if any(len(row)>width and any(v not in ['',None] for v in row[width:]) for row in grid):raise Invalid('部分数据行比表头更宽。')
                table={'headers':[str(v or '').strip() for v in grid[0][:width]],'rows':[[('' if v is None else v) for v in (row[:width]+['']*max(0,width-len(row)))] for row in grid[1:]]}
                return self.respond({'table':table,'sheets':sheets,'sheet':sheet})
            return self.respond({'error':{'message':'接口不存在。'}},404)
        except Invalid as e:return self.respond({'error':e.detail},400)
        except (ValueError,TypeError,KeyError) as e:return self.respond({'error':{'message':'数据格式无效：'+str(e)}},400)
        except Exception as e:return self.respond({'error':{'message':str(e)}},500)

class LocalHTTPServer(ThreadingHTTPServer):
    # Windows SO_REUSEADDR lets two unrelated services share the same port.
    allow_reuse_address=False
    def server_bind(self):
        if os.name=='nt':self.socket.setsockopt(socket.SOL_SOCKET,socket.SO_EXCLUSIVEADDRUSE,1)
        super().server_bind()

if __name__=='__main__':
    JOBS.mkdir(parents=True,exist_ok=True)
    try:
        local_http=urllib.request.build_opener(urllib.request.ProxyHandler({}))
        health=json.load(local_http.open(f'http://127.0.0.1:{PORT}/api/health',timeout=1))
        if health.get('app')=='vivid-workbench' and health.get('root')==str(WEB):
            if '--open' in sys.argv:webbrowser.open(f'http://127.0.0.1:{PORT}/index.html')
            sys.exit(0)
    except Exception:pass
    server=LocalHTTPServer(('127.0.0.1',PORT),Handler)
    if '--open' in sys.argv:threading.Timer(.5,lambda:webbrowser.open(f'http://127.0.0.1:{PORT}/index.html')).start()
    server.serve_forever()
