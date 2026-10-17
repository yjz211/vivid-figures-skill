"""Offline regressions for editable PPTX font selection and portable bundles.

Run with ``python tests/test_pptx_font_faces.py``. No browser, Office application,
system font installation, or source-asset modifications are required.
"""
import base64
from contextlib import nullcontext
import hashlib
import json
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest import mock
import xml.etree.ElementTree as ET


ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / 'scripts/pptx_runtime'), str(ROOT / 'architecture/lib')]

from pptx import Presentation
import install_pptx_fonts
from pptx_fonts import bundle_fonts
from pptx_text import (add_text, build_font_map, extract_embedded_fonts,
                       font_metadata, _select_font_face)
from source_asset import embed_full_fonts


FACES = [('Regular', 400, 'normal', False, False),
         ('Bold', 700, 'normal', True, False),
         ('Italic', 400, 'italic', False, True),
         ('BoldItalic', 700, 'italic', True, True)]
FONT_DIR = ROOT / 'architecture/fonts/source-assets'


class PptxFontFaceTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='vivid-pptx-font-test-')
        self.addCleanup(self.tmp.cleanup)
        self.directory = Path(self.tmp.name)
        self.svg = self.directory / 'replacement.svg'
        root = ET.Element('{http://www.w3.org/2000/svg}svg', {'viewBox': '0 0 200 200'})
        # Use the actual SVG runtime's four-face declarations, including their
        # ordering, rather than a test-only font-map approximation.
        embed_full_fonts(root, FONT_DIR, {'Arimo'})
        ET.ElementTree(root).write(self.svg, encoding='utf-8', xml_declaration=True)

    def text_item(self, weight=400, style='normal', family='Arimo'):
        return {'key': 'label', 'text': 'Editable label', 'x': 10, 'y': 30,
                'width': 90, 'fontSize': 12, 'fontFamily': family,
                'fontWeight': str(weight), 'fontStyle': style}

    def test_runtime_four_faces_select_the_matching_font_program(self):
        fonts = build_font_map(self.svg)
        self.assertEqual(len(fonts), 4)
        for face, weight, style, _, _ in FACES:
            with self.subTest(face=face):
                expected = hashlib.sha256((FONT_DIR / f'Arimo-{face}.ttf').read_bytes()).hexdigest()
                selected = _select_font_face(fonts, ['Absent Family', 'arimo'], weight, style)
                self.assertEqual(selected['sha256'], expected)

    def test_serialized_native_text_keeps_all_four_typographic_styles(self):
        fonts = build_font_map(self.svg)
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for _, weight, style, _, _ in FACES:
            add_text(slide, self.text_item(weight, style), font_map=fonts)
        output = self.directory / 'four-faces.pptx'
        prs.save(output)
        reopened = Presentation(output)
        for shape, (face, _, _, bold, italic) in zip(reopened.slides[0].shapes, FACES):
            with self.subTest(face=face):
                run = shape.text_frame.paragraphs[0].runs[0]
                self.assertEqual(run.text, 'Editable label')
                self.assertEqual(run.font.name, 'Arimo')
                self.assertEqual(run.font.bold, bold)
                self.assertEqual(run.font.italic, italic)

    def test_legacy_single_face_alias_retains_intrinsic_bold_italic(self):
        alias = 'VividSource-Arimo-BoldItalic'
        data = (FONT_DIR / 'Arimo-BoldItalic.ttf').read_bytes()
        self.svg.write_text(
            '<svg xmlns="http://www.w3.org/2000/svg"><style>'
            f'@font-face{{font-family:"{alias}";src:url(data:font/ttf;base64,'
            + base64.b64encode(data).decode('ascii') + ')}</style></svg>',
            encoding='utf-8')
        fonts = build_font_map(self.svg)
        self.assertEqual(list(fonts), [alias])
        # Caller-supplied maps from the old API have no CSS descriptor fields.
        for font_map in (fonts, {alias: font_metadata(data)}):
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            shape = add_text(slide, self.text_item(family=alias), font_map=font_map)[0]
            run = shape.text_frame.paragraphs[0].runs[0]
            self.assertEqual(run.font.name, 'Arimo')
            self.assertTrue(run.font.bold)
            self.assertTrue(run.font.italic)

    def test_extract_and_bundle_manifest_retain_every_face_and_hash(self):
        extracted = extract_embedded_fonts(self.svg, self.directory / 'extracted')
        bundle = bundle_fonts(self.svg, self.directory / 'replacement.fonts')
        manifest = json.loads((Path(bundle['directory']) / 'fonts.json').read_text(encoding='utf-8'))
        expected = {hashlib.sha256((FONT_DIR / f'Arimo-{face}.ttf').read_bytes()).hexdigest()
                    for face, *_ in FACES}
        self.assertEqual(len(extracted), 4)
        self.assertEqual(len(manifest), 4)
        self.assertEqual({info['sha256'] for info in manifest.values()}, expected)
        for info in manifest.values():
            path = Path(bundle['directory']) / info['path']
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), info['sha256'])
        mapped = build_font_map(self.svg, bundle['directory'])
        self.assertEqual({info['path'] for info in mapped.values()},
                         {str((Path(bundle['directory']) / info['path']).resolve())
                          for info in manifest.values()})

    def test_windows_installer_registers_all_faces_in_the_manifest(self):
        bundle = bundle_fonts(self.svg, self.directory / 'replacement.fonts')
        registry = mock.Mock()
        registry.CreateKey.return_value = nullcontext(object())
        registry.QueryValueEx.side_effect = FileNotFoundError
        windows = mock.Mock()
        windows.gdi32.AddFontResourceExW.return_value = 1
        fake_os = SimpleNamespace(name='nt', environ={'LOCALAPPDATA': str(self.directory / 'profile')})
        with mock.patch.dict(sys.modules, {'winreg': registry}), \
             mock.patch.object(install_pptx_fonts, 'os', fake_os), \
             mock.patch.object(install_pptx_fonts, 'ctypes', SimpleNamespace(windll=windows)):
            rows = install_pptx_fonts.install(bundle['directory'])
        self.assertEqual(len(rows), 4)
        self.assertEqual(registry.SetValueEx.call_count, 4)
        self.assertEqual(windows.gdi32.AddFontResourceExW.call_count, 4)
        self.assertEqual({' '.join(row['family'].split()) for row in rows},
                         {'Arimo', 'Arimo Bold', 'Arimo Italic', 'Arimo Bold Italic'})
        for row in rows:
            self.assertTrue(Path(row['path']).is_relative_to(self.directory))
            self.assertEqual(hashlib.sha256(Path(row['path']).read_bytes()).hexdigest(), row['sha256'])


if __name__ == '__main__':
    unittest.main()
