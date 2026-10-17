"""A single self-intersecting SVG path must retain even-odd holes in PPTX."""
from pathlib import Path
import sys
import unittest


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts/pptx_runtime'))

from pptx import Presentation
from pptx_geometry import add_geometry


def native_winding(shape, point):
    """Evaluate nonzero winding directly from the emitted DrawingML polygons."""
    winding = 0
    for path in shape._element.xpath('.//a:custGeom/a:pathLst/a:path'):
        # Path coordinates are relative to the shape and scaled to its extent.
        x = (point[0] * 12700 - shape.left) * int(path.get('w')) / shape.width
        y = (point[1] * 12700 - shape.top) * int(path.get('h')) / shape.height
        ring = []
        for command in path:
            kind = command.tag.rsplit('}', 1)[-1]
            if kind == 'moveTo':
                ring = [(int(command[0].get('x')), int(command[0].get('y')))]
            elif kind == 'lnTo':
                ring.append((int(command[0].get('x')), int(command[0].get('y'))))
            elif kind == 'close':
                for a, b in zip(ring, ring[1:] + ring[:1]):
                    side = (b[0] - a[0]) * (y - a[1]) - (x - a[0]) * (b[1] - a[1])
                    if a[1] <= y < b[1] and side > 0:
                        winding += 1
                    elif b[1] <= y < a[1] and side < 0:
                        winding -= 1
            else:
                raise AssertionError(f'Unexpected polygon command: {kind}')
    return winding


class PptxEvenOddTests(unittest.TestCase):
    def test_single_subpath_star_keeps_center_hole_and_all_five_tips(self):
        path = 'M 50 0 L 79 90 L 2 35 L 98 35 L 21 90 Z'
        for rule, center_filled in [('evenodd', False), ('nonzero', True)]:
            with self.subTest(rule=rule):
                prs = Presentation()
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                shapes = add_geometry(slide, {
                    'key': 'star', 'tag': 'path', 'attrs': {'d': path},
                    'style': {'fill': '#ff0000', 'stroke': 'none', 'fillRule': rule},
                })
                self.assertEqual(len(shapes), 1)
                shape = shapes[0]
                self.assertEqual(native_winding(shape, (50, 50)) != 0, center_filled)
                for point in [(50, 10), (15, 37), (85, 37), (27, 75), (73, 75)]:
                    self.assertNotEqual(native_winding(shape, point), 0, point)
                self.assertEqual(native_winding(shape, (5, 5)), 0)


if __name__ == '__main__':
    unittest.main()
