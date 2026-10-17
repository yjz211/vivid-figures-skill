"""Export a finished Vivid SVG as native editable PowerPoint shapes and text."""
from pathlib import Path
import argparse
import base64
import hashlib
import json
import re
import sys
import zipfile
sys.path.insert(0,str(Path(__file__).resolve().parent/'pptx_runtime'))
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx_scene import extract_scene
from pptx_geometry import add_geometry, _clip_geometry
from pptx_text import add_text, build_font_map
from pptx_groups import group_shapes


def export(svg, output, structure=None, width_pt=720, force=False, office_verify=False, install_fonts=False):
    svg=Path(svg).resolve();output=Path(output).resolve()
    if output.exists() and not force:raise FileExistsError(f'Refusing to overwrite an existing presentation: {output}')
    if width_pt<=0:raise ValueError('width-pt must be positive')
    output.parent.mkdir(parents=True,exist_ok=True)
    scene=extract_scene(svg,output.with_suffix('.svg-preview.png'))
    x,y,w,h=scene['viewBox'];scale=width_pt/w
    if max(width_pt,h*scale)>4032:raise ValueError('PowerPoint dimensions exceed 56 inches')
    prs=Presentation();prs.slide_width=Pt(width_pt);prs.slide_height=Pt(h*scale)
    slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(255,255,255)
    fonts=build_font_map(svg)
    by_key={};items={};skipped=[]
    for item in scene['items']:
        if item['tag']=='text':
            # Cropping a glyph cannot retain its normal editability. Fail explicitly.
            if item.get('clips'):
                from shapely.geometry import box
                bx,by,bw,bh=item.get('ink_bbox',item['bbox']);region=box(bx,by,bx+bw,by+bh)
                clip=_clip_geometry(item['clips'],.02/scale)
                if not clip.intersects(region):skipped.append(item['key']);continue
                if not clip.buffer(.15/scale).covers(region):raise ValueError(f'Partially clipped text {item["key"]}; remove the text clip or adjust the label before PPT export')
            shapes=add_text(slide,item,scale=scale,origin=(x,y),font_map=fonts)
        else:shapes=add_geometry(slide,item,scale=scale,origin=(x,y))
        key=item.get('source_key',item['key'])
        by_key.setdefault(key,[]).extend(shapes)
        if key not in items:items[key]=item
        elif item.get('source_key'):
            a=items[key].get('paint_bbox',items[key]['bbox']);b=item.get('paint_bbox',item['bbox'])
            xx=min(a[0],b[0]);yy=min(a[1],b[1])
            items[key]['paint_bbox']=[xx,yy,max(a[0]+a[2],b[0]+b[2])-xx,max(a[1]+a[3],b[1]+b[3])-yy]
    if structure is None:
        candidate=svg.parent/'structure.json'
        structure=candidate if candidate.exists() else None
    if isinstance(structure,(str,Path)):structure=json.loads(Path(structure).read_text(encoding='utf-8'))
    groups=group_shapes(slide,list(items.values()),by_key,structure or {})
    # Embed source keys and role names in shape names, keep all objects native.
    prs.core_properties.title=svg.stem+' — editable architecture'
    prs.core_properties.subject='Vivid native SVG geometry and editable text; connectors do not automatically reroute.'
    prs.save(output)
    from pptx_fonts import bundle_fonts, install_user_fonts
    font_bundle=bundle_fonts(svg,output.with_suffix('.fonts'))
    report={'svg':str(svg),'pptx':str(output),'svg_sha256':hashlib.sha256(svg.read_bytes()).hexdigest(),
            'size_pt':[width_pt,h*scale],'scene_items':len(scene['items']),'source_objects':len(items),
            'native_shapes':sum(len(a) for a in by_key.values()),'groups':groups,'clipped_out':skipped,
            'fonts':{name:{k:v for k,v in data.items() if k not in ('data','bytes','font_bytes')} for name,data in fonts.items()},
            'font_embedding':'not_requested','limitations':['Connections are editable paths; moving modules does not automatically reconnect arrows.','Extracted math remains editable text, not reconstructed equations.'],
            'font_bundle':font_bundle,
            'text_layout':{'editable_boxes':sum(i['tag']=='text' for i in scene['items']),
                           'reconstructed_word_spaces':sum(i.get('inferredSpaces',0) for i in scene['items'])}}
    if install_fonts:report['installed_fonts']=install_user_fonts(font_bundle['directory'])
    if office_verify:
        from pptx_office import verify_and_embed
        report['office']=verify_and_embed(output,svg)
        report['font_embedding']=report['office']['font_embedding']
    with zipfile.ZipFile(output) as z:
        report['media_files']=[n for n in z.namelist() if n.startswith('ppt/media/')]
    output.with_suffix('.pptx-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    return report


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('svg',type=Path);parser.add_argument('--out',required=True,type=Path)
    parser.add_argument('--structure',type=Path);parser.add_argument('--width-pt',type=float,default=720)
    parser.add_argument('--force',action='store_true');parser.add_argument('--office-verify',action='store_true',help='Windows PowerPoint: embed fonts and render an actual slide preview')
    parser.add_argument('--install-fonts',action='store_true',help='Install the bundled fonts for the current Windows user')
    args=parser.parse_args()
    result=export(args.svg,args.out,args.structure,args.width_pt,args.force,args.office_verify,args.install_fonts)
    print(json.dumps({k:result[k] for k in ['pptx','native_shapes','font_embedding','media_files']},ensure_ascii=False))


if __name__=='__main__':main()
