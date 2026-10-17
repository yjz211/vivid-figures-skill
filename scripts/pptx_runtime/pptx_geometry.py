"""Native editable SVG geometry for PowerPoint. No FigGenie runtime dependency.

Path parser/arc mathematics and DrawingML construction adapted from svg2pptx.py,
Copyright (c) 2026 Sleepy Avacado, MIT; see NOTICE-pptx.txt.
All other conversion, clipping, stroke and computed-style logic authored for Vivid.
"""
import math
import re
from lxml import etree
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

def mapply(m,x,y):
    return m[0]*x+m[2]*y+m[4], m[1]*x+m[3]*y+m[5]

_NUM = re.compile(r'[-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?')


def arc_to_cubics(x1, y1, rx, ry, phi, fa, fs, x2, y2):
    """SVG endpoint arc → list of cubic segments (c1, c2, p) per SVG spec F.6.5."""
    if rx == 0 or ry == 0 or (x1 == x2 and y1 == y2):
        return [((x1, y1), (x2, y2), (x2, y2))]
    phi = math.radians(phi); cp, sp = math.cos(phi), math.sin(phi)
    dx, dy = (x1 - x2) / 2, (y1 - y2) / 2
    x1p, y1p = cp * dx + sp * dy, -sp * dx + cp * dy
    rx, ry = abs(rx), abs(ry)
    lam = (x1p ** 2) / (rx ** 2) + (y1p ** 2) / (ry ** 2)
    if lam > 1: rx *= math.sqrt(lam); ry *= math.sqrt(lam)
    num = rx ** 2 * ry ** 2 - rx ** 2 * y1p ** 2 - ry ** 2 * x1p ** 2
    den = rx ** 2 * y1p ** 2 + ry ** 2 * x1p ** 2
    coef = (1 if fa != fs else -1) * math.sqrt(max(0.0, num / den)) if den else 0.0
    cxp, cyp = coef * rx * y1p / ry, -coef * ry * x1p / rx
    cx, cy = cp * cxp - sp * cyp + (x1 + x2) / 2, sp * cxp + cp * cyp + (y1 + y2) / 2

    def ang(ux, uy, vx, vy):
        d = math.atan2(ux * vy - uy * vx, ux * vx + uy * vy)
        return d
    t1 = ang(1, 0, (x1p - cxp) / rx, (y1p - cyp) / ry)
    dt = ang((x1p - cxp) / rx, (y1p - cyp) / ry, (-x1p - cxp) / rx, (-y1p - cyp) / ry)
    if not fs and dt > 0: dt -= 2 * math.pi
    if fs and dt < 0: dt += 2 * math.pi
    n = max(1, int(math.ceil(abs(dt) / (math.pi / 2) - 1e-9)))
    out, a0 = [], t1
    for i in range(n):
        a1 = t1 + dt * (i + 1) / n
        k = 4 / 3 * math.tan((a1 - a0) / 4)

        def pt(a):
            return (cx + rx * math.cos(a) * cp - ry * math.sin(a) * sp, cy + rx * math.cos(a) * sp + ry * math.sin(a) * cp)

        def dpt(a):
            return (-rx * math.sin(a) * cp - ry * math.cos(a) * sp, -rx * math.sin(a) * sp + ry * math.cos(a) * cp)
        p0, p3 = pt(a0), pt(a1); d0, d3 = dpt(a0), dpt(a1)
        out.append(((p0[0] + k * d0[0], p0[1] + k * d0[1]), (p3[0] - k * d3[0], p3[1] - k * d3[1]), p3))
        a0 = a1
    return out


def parse_path(d):
    """→ list of subpaths: {'pts': [('L', x, y) | ('C', x1, y1, x2, y2, x, y)], 'start': (x, y), 'closed': bool}"""
    toks = re.findall(r'[MmZzLlHhVvCcSsQqTtAa]|' + _NUM.pattern, d or '')
    subs, cur, i = [], None, 0
    x = y = sx = sy = 0.0; lc = None; cmd = None
    while i < len(toks):
        t = toks[i]
        if re.match(r'[A-Za-z]', t):
            cmd = t; i += 1
            if cmd in 'Zz':
                if cur: cur['closed'] = True; x, y = sx, sy
                lc = None
                continue
        else:
            if cmd is None: break
            if cmd == 'M': cmd = 'L'
            elif cmd == 'm': cmd = 'l'

        def num(k):
            return [float(v) for v in toks[i:i + k]]
        if cmd in 'Mm':
            px, py = num(2); i += 2
            if cmd == 'm': px += x; py += y
            x, y = sx, sy = px, py
            cur = {'start': (x, y), 'pts': [], 'closed': False}; subs.append(cur); lc = None
        elif cmd in 'Ll':
            px, py = num(2); i += 2
            if cmd == 'l': px += x; py += y
            x, y = px, py; cur['pts'].append(('L', x, y)); lc = None
        elif cmd in 'Hh':
            px = num(1)[0]; i += 1
            x = px + x if cmd == 'h' else px; cur['pts'].append(('L', x, y)); lc = None
        elif cmd in 'Vv':
            py = num(1)[0]; i += 1
            y = py + y if cmd == 'v' else py; cur['pts'].append(('L', x, y)); lc = None
        elif cmd in 'Cc':
            v = num(6); i += 6
            if cmd == 'c': v = [v[0] + x, v[1] + y, v[2] + x, v[3] + y, v[4] + x, v[5] + y]
            cur['pts'].append(('C', *v)); lc = (v[2], v[3]); x, y = v[4], v[5]
        elif cmd in 'Ss':
            v = num(4); i += 4
            if cmd == 's': v = [v[0] + x, v[1] + y, v[2] + x, v[3] + y]
            c1 = (2 * x - lc[0], 2 * y - lc[1]) if lc and lc[0] != 'Q' else (x, y)
            cur['pts'].append(('C', c1[0], c1[1], v[0], v[1], v[2], v[3])); lc = (v[0], v[1]); x, y = v[2], v[3]
        elif cmd in 'Qq':
            v = num(4); i += 4
            if cmd == 'q': v = [v[0] + x, v[1] + y, v[2] + x, v[3] + y]
            c1 = (x + 2 / 3 * (v[0] - x), y + 2 / 3 * (v[1] - y)); c2 = (v[2] + 2 / 3 * (v[0] - v[2]), v[3] + 2 / 3 * (v[1] - v[3]))
            cur['pts'].append(('C', c1[0], c1[1], c2[0], c2[1], v[2], v[3])); lc = ('Q', v[0], v[1]); x, y = v[2], v[3]
        elif cmd in 'Tt':
            v = num(2); i += 2
            if cmd == 't': v = [v[0] + x, v[1] + y]
            qc = (2 * x - lc[1], 2 * y - lc[2]) if (lc and lc[0] == 'Q') else (x, y)
            c1 = (x + 2 / 3 * (qc[0] - x), y + 2 / 3 * (qc[1] - y)); c2 = (v[0] + 2 / 3 * (qc[0] - v[0]), v[1] + 2 / 3 * (qc[1] - v[1]))
            cur['pts'].append(('C', c1[0], c1[1], c2[0], c2[1], v[0], v[1])); lc = ('Q', qc[0], qc[1]); x, y = v[0], v[1]
        elif cmd in 'Aa':
            v = num(7); i += 7
            ex, ey = (v[5] + x, v[6] + y) if cmd == 'a' else (v[5], v[6])
            for c1, c2, p in arc_to_cubics(x, y, v[0], v[1], v[2], int(v[3]), int(v[4]), ex, ey):
                cur['pts'].append(('C', c1[0], c1[1], c2[0], c2[1], p[0], p[1]))
            x, y = ex, ey; lc = None
        else:
            break
    return subs


def transform_subpaths(subs, m):
    out = []
    for s in subs:
        pts = []
        for p in s['pts']:
            if p[0] == 'L':
                pts.append(('L', *mapply(m, p[1], p[2])))
            else:
                pts.append(('C', *mapply(m, p[1], p[2]), *mapply(m, p[3], p[4]), *mapply(m, p[5], p[6])))
        out.append({'start': mapply(m, *s['start']), 'pts': pts, 'closed': s['closed']})
    return out


class GeometryError(ValueError):
    """Unsupported or ambiguous geometry must not become a silent approximation."""


CURVE_TOLERANCE_PT = 0.02


def _num(v, default=0):
    if v is None or v == '':
        return float(default)
    try:
        return float(v)
    except (TypeError, ValueError):
        m = re.fullmatch(r'\s*(' + _NUM.pattern + r')(?:px)?\s*', str(v))
        if not m:
            raise GeometryError(f'Expected resolved SVG number, got {v!r}')
        return float(m.group(1))


def _primitive(item):
    a, tag = item.get('attrs', {}), item['tag']
    n = lambda k, default=0: _num(a.get(k), default)
    if tag == 'path':
        return parse_path(a.get('d', ''))
    if tag == 'line':
        return [{'start': (n('x1'), n('y1')), 'pts': [('L', n('x2'), n('y2'))], 'closed': False}]
    if tag in ('polyline', 'polygon'):
        vals = [float(v) for v in _NUM.findall(a.get('points', ''))]
        if len(vals) % 2:
            raise GeometryError('Odd polygon coordinate count')
        points = list(zip(vals[::2], vals[1::2]))
        return [{'start': points[0], 'pts': [('L', *p) for p in points[1:]], 'closed': tag == 'polygon'}] if points else []
    if tag == 'rect':
        x, y, w, h = n('x'), n('y'), n('width'), n('height')
        if w <= 0 or h <= 0:
            return []
        rx = min(w/2, max(0, n('rx', n('ry'))))
        ry = min(h/2, max(0, n('ry', n('rx'))))
        if rx and ry:
            return parse_path(f'M{x+rx},{y} H{x+w-rx} A{rx},{ry} 0 0 1 {x+w},{y+ry} V{y+h-ry} A{rx},{ry} 0 0 1 {x+w-rx},{y+h} H{x+rx} A{rx},{ry} 0 0 1 {x},{y+h-ry} V{y+ry} A{rx},{ry} 0 0 1 {x+rx},{y} Z')
        return parse_path(f'M{x},{y} H{x+w} V{y+h} H{x} Z')
    if tag in ('circle', 'ellipse'):
        x, y = n('cx'), n('cy')
        rx, ry = (n('r'), n('r')) if tag == 'circle' else (n('rx'), n('ry'))
        if rx <= 0 or ry <= 0:
            return []
        return parse_path(f'M{x+rx},{y} A{rx},{ry} 0 0 1 {x-rx},{y} A{rx},{ry} 0 0 1 {x+rx},{y} Z')
    raise GeometryError(f'Unsupported geometry tag: {tag}')


def _color(value, opacity=1):
    from PIL import ImageColor
    value = str(value or 'none').strip()
    if value in ('none', 'transparent'):
        return None
    if value.startswith('url(') or value == 'currentColor':
        raise GeometryError(f'Paint needs resolved solid color: {value}')
    try:
        rgba = re.fullmatch(r'rgba?\(([^)]+)\)', value)
        if rgba:
            parts = re.split(r'\s*[,/]\s*|\s+', rgba.group(1).strip())
            rgb = [round(float(x[:-1])*2.55) if x.endswith('%') else round(float(x)) for x in parts[:3]]
            alpha = float(parts[3].rstrip('%')) / (100 if parts[3].endswith('%') else 1) if len(parts) > 3 else 1
        else:
            r, g, b, a = ImageColor.getcolor(value, 'RGBA')
            rgb, alpha = [r, g, b], a/255
        return ''.join(f'{min(255,max(0,c)):02X}' for c in rgb), max(0, min(1, float(opacity)*alpha))
    except (ValueError, TypeError) as exc:
        raise GeometryError(f'Unresolved color {value!r}') from exc


def _solid(parent, paint):
    if paint is None:
        etree.SubElement(parent, qn('a:noFill'))
        return
    solid = etree.SubElement(parent, qn('a:solidFill'))
    rgb = etree.SubElement(solid, qn('a:srgbClr'), val=paint[0])
    if paint[1] < 1:
        etree.SubElement(rgb, qn('a:alpha'), val=str(round(paint[1]*100000)))


def _shape(slide, subs, name, fill, stroke=None, width=1, style=None):
    """Write one compound native DrawingML path, preserving cubic control points."""
    if not subs:
        return None
    xy = [s['start'] for s in subs]
    for sub in subs:
        for seg in sub['pts']:
            xy.extend(zip(seg[1::2], seg[2::2]))
    x0, y0 = min(p[0] for p in xy), min(p[1] for p in xy)
    w, h = max(max(p[0] for p in xy)-x0, .0001), max(max(p[1] for p in xy)-y0, .0001)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x0), Pt(y0), Pt(w), Pt(h))
    shape.name = str(name)
    sp = shape._element.spPr
    for child in list(sp):
        if child.tag != qn('a:xfrm'):
            sp.remove(child)
    geom = etree.SubElement(sp, qn('a:custGeom'))
    for tag in ('avLst', 'gdLst', 'ahLst', 'cxnLst'):
        etree.SubElement(geom, qn('a:'+tag))
    etree.SubElement(geom, qn('a:rect'), l='0', t='0', r='r', b='b')
    pathlist = etree.SubElement(geom, qn('a:pathLst'))
    path = etree.SubElement(pathlist, qn('a:path'), w=str(max(1, round(w*12700))), h=str(max(1, round(h*12700))), fill='norm' if fill else 'none', stroke='1' if stroke else '0')
    def point(el, x, y):
        etree.SubElement(el, qn('a:pt'), x=str(round((x-x0)*12700)), y=str(round((y-y0)*12700)))
    for sub in subs:
        point(etree.SubElement(path, qn('a:moveTo')), *sub['start'])
        for seg in sub['pts']:
            el = etree.SubElement(path, qn('a:lnTo' if seg[0] == 'L' else 'a:cubicBezTo'))
            for x, y in zip(seg[1::2], seg[2::2]):
                point(el, x, y)
        if sub['closed']:
            etree.SubElement(path, qn('a:close'))
    _solid(sp, fill)
    line = etree.SubElement(sp, qn('a:ln'), w=str(round(max(0, width)*12700)))
    _solid(line, stroke)
    if stroke:
        style = style or {}
        line.set('cap', {'round':'rnd','square':'sq'}.get(style.get('linecap'), 'flat'))
        # Dash lengths are relative to actual local stroke width; CTM cancels.
        dash = _dashes(style.get('dasharray'))
        if dash:
            cd = etree.SubElement(line, qn('a:custDash'))
            sw = _num(style.get('strokeWidth'), 1)
            for d, s in zip(dash[::2], dash[1::2]):
                etree.SubElement(cd, qn('a:ds'), d=str(round(d/max(sw,1e-12)*100000)), sp=str(round(s/max(sw,1e-12)*100000)))
        join = style.get('linejoin', 'miter')
        etree.SubElement(line, qn('a:'+({'round':'round','bevel':'bevel'}.get(join,'miter'))))
    etree.SubElement(sp, qn('a:effectLst'))
    shape.text_frame.clear()
    return shape


def _flatten(subs, tolerance):
    def cubic(p0, p1, p2, p3, depth=0):
        # Distance to chord plus excess polygon length catches folded straight cubics.
        chord = math.dist(p0, p3)
        excess = math.dist(p0,p1)+math.dist(p1,p2)+math.dist(p2,p3)-chord
        def distance(p):
            if not chord:
                return math.dist(p,p0)
            return abs((p3[0]-p0[0])*(p0[1]-p[1])-(p0[0]-p[0])*(p3[1]-p0[1]))/chord
        if depth >= 22 or max(distance(p1),distance(p2),excess) <= tolerance:
            return [p3]
        mid = lambda a,b: ((a[0]+b[0])/2,(a[1]+b[1])/2)
        a,b,c = mid(p0,p1),mid(p1,p2),mid(p2,p3)
        d,e = mid(a,b),mid(b,c)
        f = mid(d,e)
        return cubic(p0,a,d,f,depth+1)+cubic(f,e,c,p3,depth+1)
    rings = []
    for s in subs:
        pts=[s['start']]
        for seg in s['pts']:
            if seg[0] == 'L':
                pts.append((seg[1],seg[2]))
            else:
                pts.extend(cubic(pts[-1],(seg[1],seg[2]),(seg[3],seg[4]),(seg[5],seg[6])))
        if s['closed'] and pts[-1] != pts[0]:
            pts.append(pts[0])
        rings.append(pts)
    return rings


def _filled(subs, rule, tolerance):
    from shapely.geometry import LineString, GeometryCollection
    from shapely.ops import polygonize, unary_union
    rings = [r if r[0] == r[-1] else r+[r[0]] for r in _flatten(subs,tolerance) if len(r)>2]
    if not rings:
        return GeometryCollection()
    cells = list(polygonize(unary_union([LineString(r) for r in rings])))
    def winding(p):
        total=0
        for ring in rings:
            for a,b in zip(ring,ring[1:]):
                side=(b[0]-a[0])*(p.y-a[1])-(p.x-a[0])*(b[1]-a[1])
                if a[1] <= p.y < b[1] and side>0: total+=1
                if b[1] <= p.y < a[1] and side<0: total-=1
        return total
    return unary_union([p for p in cells if (abs(winding(p.representative_point()))%2 if rule == 'evenodd' else winding(p.representative_point())!=0)])


def _polygon_subs(geometry):
    from shapely.geometry.polygon import orient
    out=[]
    def add(g):
        if g.is_empty:
            return
        if g.geom_type == 'Polygon':
            g=orient(g,sign=1)
            for ring in [g.exterior,*g.interiors]:
                pts=list(ring.coords)
                out.append({'start':pts[0],'pts':[('L',*p) for p in pts[1:-1]],'closed':True})
        elif hasattr(g,'geoms'):
            for child in g.geoms: add(child)
    add(geometry)
    return out


def _dashes(value):
    if value is None or value == '' or value == 'none':
        return []
    values = [_num(x) for x in value] if isinstance(value,(list,tuple)) else [_num(x) for x in re.split(r'[,\s]+',str(value).strip())]
    if any(x<0 for x in values):
        raise GeometryError('Negative stroke dash')
    if not sum(values):
        return []
    return values*2 if len(values)%2 else values


def _stroke_geometry(subs, style, tolerance):
    from shapely.geometry import LineString, Point, GeometryCollection
    from shapely.ops import unary_union, substring
    width = _num(style.get('strokeWidth'),1)
    if width <= 0:
        return GeometryCollection()
    dash = _dashes(style.get('dasharray'))
    offset = _num(style.get('dashoffset'))
    cap={'butt':2,'round':1,'square':3}.get(style.get('linecap'),2)
    join={'miter':2,'round':1,'bevel':3}.get(style.get('linejoin'),2)
    output=[]
    # Quarter circle segment count keeps buffer approximation within print tolerance.
    q=max(8, int(math.ceil(math.pi/(4*math.acos(max(-1,1-min(tolerance,width/2)/(width/2)))))))
    for pts in _flatten(subs,tolerance):
        if len(pts)<2: continue
        line=LineString(pts)
        if not dash:
            output.append(line.buffer(width/2,quad_segs=q,cap_style=cap,join_style=join,mitre_limit=_num(style.get('miterlimit'),4)))
            continue
        pos=-offset%sum(dash)
        if pos>0: pos-=sum(dash)
        idx=0
        while pos<line.length:
            length=dash[idx%len(dash)]
            if idx%2==0 and pos+length>=0:
                piece=substring(line,max(pos,0),min(pos+length,line.length))
                output.append(piece.buffer(width/2,quad_segs=q,cap_style=cap,join_style=join,mitre_limit=_num(style.get('miterlimit'),4)))
            pos+=length
            idx+=1
    return unary_union(output)


def _clip_geometry(clips, tolerance):
    from shapely.geometry import box
    from shapely.ops import unary_union
    result=None
    for clip in clips:
        if not isinstance(clip,dict):
            raise GeometryError('Clip must be a global-viewBox path or rect record')
        if 'children' in clip:
            g=unary_union([_clip_geometry([child],tolerance) for child in clip['children']])
        elif 'rect' in clip:
            x,y,w,h=clip['rect']; g=box(x,y,x+w,y+h)
        elif 'd' in clip:
            g=_filled(parse_path(clip['d']),clip.get('fillRule','nonzero'),tolerance)
        else:
            raise GeometryError('Unresolved clip; expected d or rect in global viewBox units')
        result=g if result is None else result.intersection(g)
    return result


def add_geometry(slide, item, scale=1, origin=(0,0)):
    """Append native geometry in paint order; return all created shapes.

    matrix maps local coordinates to viewBox. Final position is
    (viewBox - origin) * scale, in PowerPoint points. Clips are intersections
    in global viewBox coordinates. Curves used for clipping/stroke outlines
    are flattened to 0.02 print-point tolerance (not rasterized).
    """
    if scale<=0:
        raise GeometryError('scale must be positive')
    subs=_primitive(item)
    if not subs:
        return []
    m=list(item.get('matrix',[1,0,0,1,0,0]))
    st=dict(item.get('style',{}))
    if st.get('vectorEffect','none') not in ('none',''):
        raise GeometryError('Non-scaling stroke needs an explicit scene-space stroke conversion')
    opacity=_num(st.get('opacity'),1)
    fill=_color(st.get('fill','black'),opacity*_num(st.get('fillOpacity'),1))
    stroke=_color(st.get('stroke','none'),opacity*_num(st.get('strokeOpacity'),1))
    width=_num(st.get('strokeWidth'),1)
    if width<=0: stroke=None
    if not fill and not stroke: return []
    if fill and fill[1]<=0: fill=None
    if stroke and stroke[1]<=0: stroke=None
    sx,sy=math.hypot(m[0],m[1]),math.hypot(m[2],m[3])
    if not sx or not sy:
        return []
    uniform=abs(sx-sy)<1e-7*max(sx,sy) and abs(m[0]*m[2]+m[1]*m[3])<1e-7*sx*sy
    global_subs=transform_subpaths(subs,m)
    final_m=[scale,0,0,scale,-origin[0]*scale,-origin[1]*scale]
    key=str(item.get('key','geometry'))
    clips=item.get('clips') or []
    rule=st.get('fillRule','nonzero')
    if rule not in ('evenodd','nonzero'):
        raise GeometryError(f'Unknown fill rule {rule}')
    clip=None
    if clips:
        from shapely.geometry import box
        clip=_clip_geometry(clips,CURVE_TOLERANCE_PT/scale)
        # Drop redundant viewport clips while retaining the original cubic path.
        pts=[s['start'] for s in global_subs]
        for sub in global_subs:
            for seg in sub['pts']: pts.extend(zip(seg[1::2],seg[2::2]))
        pad=(width/2*(sx+sy)*max(1,_num(st.get('miterlimit'),4))) if stroke else 0
        bbox=box(min(p[0] for p in pts)-pad,min(p[1] for p in pts)-pad,max(p[0] for p in pts)+pad,max(p[1] for p in pts)+pad)
        if clip.covers(bbox):
            clips=[]
            clip=None
    # DrawingML uses nonzero winding. Even one self-intersecting subpath can
    # differ under even-odd filling, so normalize every visible even-odd fill.
    special=bool(clips) or (stroke and (not uniform or _num(st.get('dashoffset'))!=0)) or (fill and rule=='evenodd')
    order=st.get('paintOrder',st.get('paint-order',item.get('attrs',{}).get('paint-order','normal')))
    stroke_first=str(order).strip().startswith('stroke')
    if not special:
        final=transform_subpaths(global_subs,final_m)
        if stroke_first and fill and stroke:
            return [_shape(slide,final,key+'__stroke',None,stroke,width*sx*scale,st),_shape(slide,final,key,fill)]
        return [_shape(slide,final,key,fill,stroke,width*sx*scale,st)]
    from shapely.affinity import affine_transform
    tolerance=CURVE_TOLERANCE_PT/scale
    jobs=[]
    if fill:
        g=_filled(global_subs,rule,tolerance)
        if clip is not None: g=g.intersection(clip)
        jobs.append(('fill',g,fill))
    if stroke:
        # Buffer in local space before CTM, preserving nonuniform scale/shear.
        g=_stroke_geometry(subs,st,tolerance/max(sx+sy,1e-8))
        g=affine_transform(g,[m[0],m[2],m[1],m[3],m[4],m[5]])
        if clip is not None: g=g.intersection(clip)
        jobs.append(('stroke',g,stroke))
    if stroke_first: jobs.reverse()
    result=[]
    for kind,g,paint in jobs:
        if g.is_empty: continue
        final=transform_subpaths(_polygon_subs(g),final_m)
        shape=_shape(slide,final,key if len(jobs)==1 else key+'__'+kind,paint)
        if shape is not None: result.append(shape)
    return result
